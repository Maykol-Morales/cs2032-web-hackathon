#!/usr/bin/env python3
"""Genera src/data/results.json cruzando inscripciones, entregas y certificados."""
import csv
import json
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC_DATA = ROOT / "src" / "data"
SRC_DATA.mkdir(parents=True, exist_ok=True)

# Datos privados de participantes (ignorados por git)
PRIVATE_DATA = ROOT / "data"

INSCRIPCIONES_CSV = PRIVATE_DATA / (
    "Formulario de inscripción HACK __ UTEC (respuestas) - Respuestas de formulario 1.csv"
)
ENTREGAS_CSV = PRIVATE_DATA / (
    "Entregas HACK__UTEC Cloud Computing (Respuestas) - Respuestas de formulario 1.csv"
)
CERTIFICADOS_XLSX = PRIVATE_DATA / "certificados.xlsx"


def repair_excel_text(s: str) -> str:
    """Repara cadenas UTF-8 mal codificadas como latin1 que vienen del xlsx."""
    if s is None:
        return ""
    s = str(s)
    # OpenPyXL escapa caracteres de control como _x0081_; los convertimos de vuelta a bytes.
    def unescape_control(match: re.Match) -> str:
        code = int(match.group(1), 16)
        return chr(code)
    s = re.sub(r"_x([0-9a-fA-F]{4})_", unescape_control, s)
    try:
        s = s.encode("latin1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return s


def normalize_text(s: str) -> str:
    """Limpia y normaliza una cadena para comparaciones."""
    if s is None:
        return ""
    s = str(s).strip()
    s = repair_excel_text(s)
    s = s.lower()
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    s = re.sub(r"[^a-z0-9]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def normalize_team_name(name: str) -> str:
    """Normaliza un nombre de equipo para hacer coincidir inscripciones/entregas."""
    if not name:
        return ""
    n = normalize_text(name)
    replacements = {
        "claude y dos mas": "claude y 3 mas",
        "claude y 3 mas": "claude y 3 mas",
        "mapfre": "seguros mapfre",
        "segurosmapfre": "seguros mapfre",
        "momazos cloud comp": "momazos cloud comp",
        "momazoscloud comp": "momazos cloud comp",
        "momazoscloudcomp": "momazos cloud comp",
        "fvpf fabiana vuelve por favor": "fvpf",
        "los rimacsitos rafael choque sebastian loli y andre": "los rimacsitos",
        "ref 1 lexion": "ggno team",
        "ref1lexion": "ggno team",
        "gg no team": "ggno team",
        "ggnoteam": "ggno team",
    }
    return replacements.get(n, n)


def load_inscripciones() -> list[dict]:
    """Cada fila del CSV de inscripciones es un registro separado."""
    records: list[dict] = []
    seen_keys: set[str] = set()
    with open(INSCRIPCIONES_CSV, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            raw_name = row.get("Nombre de equipo", "").strip()
            if not raw_name:
                continue
            if raw_name.lower() == "test":
                continue
            email = row.get("Dirección de correo electrónico", "").strip().lower()
            members = []
            for i in range(1, 4):
                member_name = row.get(f"Nombre completo N°{i}", "").strip()
                student_code = row.get(f"Código de alumno N°{i}", "").strip()
                section = row.get(f"Sección N°{i}", "").strip()
                if member_name:
                    members.append(
                        {
                            "name": member_name,
                            "code": student_code,
                            "section": section,
                        }
                    )
            # Deduplicar miembros por nombre normalizado
            seen_member_names: set[str] = set()
            unique_members = []
            for m in members:
                key = normalize_text(m["name"])
                if key and key not in seen_member_names:
                    seen_member_names.add(key)
                    unique_members.append(m)
            records.append(
                {
                    "name": raw_name,
                    "normalizedName": normalize_team_name(raw_name),
                    "members": unique_members,
                    "emails": [email] if email else [],
                }
            )
            seen_keys.add(raw_name.lower())
    return records


def load_entregas() -> list[dict]:
    submissions: list[dict] = []
    with open(ENTREGAS_CSV, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            raw_name = row.get("Nombre de equipo", "").strip()
            if not raw_name:
                continue
            try:
                score = float(row.get("RESULTADO", "0") or 0)
            except ValueError:
                score = 0.0
            submissions.append(
                {
                    "name": raw_name,
                    "normalizedName": normalize_team_name(raw_name),
                    "score": score,
                    "github": row.get("GitHub", "").strip(),
                    "solution": row.get("Solución", "").strip(),
                    "email": row.get("Dirección de correo electrónico", "").strip().lower(),
                }
            )
    return submissions


def load_certificados() -> dict[str, dict]:
    certs: dict[str, dict] = {}
    wb = load_workbook(CERTIFICADOS_XLSX, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue
        raw_name, position, url = row[0], row[1], row[2]
        if not raw_name or not url:
            continue
        name = repair_excel_text(str(raw_name))
        key = normalize_text(name)
        certs[key] = {
            "name": name,
            "url": str(url).strip(),
            "position": position,
        }
    return certs


def find_best_team_for_submission(sub: dict, teams: list[dict]) -> dict | None:
    """Encuentra el equipo inscrito que mejor coincida con una entrega."""
    # 1. Coincidencia por nombre normalizado
    name_matches = [t for t in teams if t["normalizedName"] == sub["normalizedName"]]
    if len(name_matches) == 1:
        return name_matches[0]
    if len(name_matches) > 1:
        # Elegir el que tenga más miembros
        return max(name_matches, key=lambda t: len(t["members"]))

    # 2. Coincidencia por correo del entregante
    if sub["email"]:
        email_matches = [t for t in teams if sub["email"] in t["emails"]]
        if len(email_matches) == 1:
            return email_matches[0]
        if len(email_matches) > 1:
            # Si hay varios, preferir el que tenga más miembros
            return max(email_matches, key=lambda t: len(t["members"]))

    return None


def main() -> None:
    teams = load_inscripciones()
    submissions = load_entregas()
    certs = load_certificados()

    # Cruzar entregas con equipos inscritos
    matched_teams: dict[str, dict] = {}
    unmatched_submissions: list[dict] = []

    for sub in submissions:
        team = find_best_team_for_submission(sub, teams)
        if team is None:
            unmatched_submissions.append(sub)
            continue

        # Clave estable para no duplicar si el mismo equipo entregó varias veces
        team_key = normalize_team_name(team["name"])
        if team_key not in matched_teams:
            matched_teams[team_key] = {
                "key": team_key,
                "name": team["name"],
                "score": sub["score"],
                "github": sub["github"],
                "solution": sub["solution"],
                "members": team["members"],
            }
        else:
            existing = matched_teams[team_key]
            if sub["score"] > existing["score"]:
                existing["score"] = sub["score"]
                existing["github"] = sub["github"]
                existing["solution"] = sub["solution"]
            # Fusionar miembros si hay diferencias
            seen = {normalize_text(m["name"]) for m in existing["members"]}
            for m in team["members"]:
                key = normalize_text(m["name"])
                if key and key not in seen:
                    seen.add(key)
                    existing["members"].append(m)

    # Asignar certificados a cada miembro
    for team in matched_teams.values():
        for member in team["members"]:
            key = normalize_text(member["name"])
            if key in certs:
                member["certificateUrl"] = certs[key]["url"]
                member["certificateFound"] = True
            else:
                member["certificateFound"] = False
                # Matching parcial como fallback
                for cert_key, cert in certs.items():
                    if key and cert_key and (key in cert_key or cert_key in key):
                        member["certificateUrl"] = cert["url"]
                        member["certificateFound"] = True
                        break

    participated = sorted(matched_teams.values(), key=lambda t: (-t["score"], t["name"]))

    # Determinar podio usando el xlsx como fuente principal de puestos
    podium: list[dict] = []
    podium_keys: set[str] = set()

    cert_positions: dict[int, list[str]] = {1: [], 2: [], 3: []}
    for cert in certs.values():
        pos = cert.get("position")
        if isinstance(pos, (int, float)) and 1 <= pos <= 3:
            cert_positions[int(pos)].append(cert["name"])

    for pos in (1, 2, 3):
        for cert_name in cert_positions.get(pos, []):
            cert_key = normalize_text(cert_name)
            for team in participated:
                for member in team["members"]:
                    if normalize_text(member["name"]) == cert_key:
                        if team["key"] not in podium_keys:
                            podium_keys.add(team["key"])
                            podium.append({"position": pos, **team})
                        break
                if team["key"] in podium_keys:
                    break

    # Completar con mejores puntajes si faltan puestos
    for team in participated:
        if len(podium) >= 3:
            break
        if team["key"] not in podium_keys:
            next_pos = len(podium) + 1
            podium_keys.add(team["key"])
            podium.append({"position": next_pos, **team})

    podium.sort(key=lambda t: t["position"])
    non_podium = [t for t in participated if t["key"] not in podium_keys]

    result = {
        "generatedAt": "2026-06-30",
        "podium": podium,
        "teams": non_podium,
        "unmatchedSubmissions": [
            {"name": s["name"], "score": s["score"]} for s in unmatched_submissions
        ],
    }

    out_path = SRC_DATA / "results.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"Escrito: {out_path}")
    print(f"Equipos con datos completos: {len(participated)}")
    print(f"Podio: {[f'{p['position']}. {p['name']} ({p['score']})' for p in podium]}")
    if unmatched_submissions:
        print("Entregas sin equipo inscrito:")
        for s in unmatched_submissions:
            print(f"  - {s['name']} score={s['score']}")


if __name__ == "__main__":
    main()
